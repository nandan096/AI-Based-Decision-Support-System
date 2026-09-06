import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def create_document():
    doc = docx.Document()

    # --- Page Setup (A4, IEEE Margins) ---
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # --- Base Style Settings ---
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # Helper function for body paragraphs
    def add_p(text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=4, space_before=0, line_spacing=1.05, bold=False, italic=False):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.line_spacing = line_spacing
        if text:
            run = p.add_run(text)
            run.bold = bold
            run.italic = italic
            run.font.name = 'Times New Roman'
        return p

    # Helper function for Section Headings
    def add_heading(num_and_title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(num_and_title)
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'
        return p

    # Helper function for Subsections (A., B., C.)
    def add_subheading(sub_title):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(sub_title)
        run.italic = True
        run.bold = True
        run.font.size = Pt(10)
        run.font.name = 'Times New Roman'
        return p

    # ==================== TITLE ====================
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run("AI-Based Decision Support System for Agricultural Price Prediction and Market Recommendation")
    run_title.font.name = 'Times New Roman'
    run_title.font.size = Pt(20)
    run_title.bold = True

    # ==================== AUTHOR BLOCK TABLE ====================
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    authors = [
        [
            ("Nandakumar G", True, Pt(11)),
            ("Department of MCA", False, Pt(10)),
            ("Dayananda Sagar College of Engineering", False, Pt(10)),
            ("Bengaluru, India", False, Pt(10)),
            ("nandakumargs05@gmail.com", False, Pt(9.5))
        ],
        [
            ("Prof. Mahendra Kumar", True, Pt(11)),
            ("Department of MCA", False, Pt(10)),
            ("Dayananda Sagar College of Engineering", False, Pt(10)),
            ("Bengaluru, India", False, Pt(10)),
            ("mahendra-mcavtu@dayanandasagar.edu", False, Pt(9.5))
        ]
    ]

    for col_idx, author_lines in enumerate(authors):
        cell = table.cell(0, col_idx)
        cell.width = Inches(3.4)
        for i, (text, is_bold, font_sz) in enumerate(author_lines):
            p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(text)
            r.bold = is_bold
            r.font.size = font_sz
            r.font.name = 'Times New Roman'

    # Remove table borders
    for row in table.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(
                f'<w:tcBorders {nsdecls("w")}>'
                f'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
                f'</w:tcBorders>'
            )
            tcPr.append(tcBorders)

    # Empty spacer
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(8)

    # ==================== ABSTRACT & KEYWORDS ====================
    p_abs = add_p()
    r_abs_label = p_abs.add_run("Abstract— ")
    r_abs_label.bold = True
    r_abs_label.italic = True
    r_abs_body = p_abs.add_run(
        "Spot markets for agricultural commodities in developing nations are characterized by severe volatility in "
        "time and prices in space, which exposes small-holder farmers to asymmetric risks and distress selling. Although modern "
        "studies have investigated statistical as well as machine learning techniques for price prediction, current implementations "
        "tend to be divorced between prediction modeling and post-harvest economic decision-making. This paper proposes a Decision "
        "Support System (DSS) that is built based on a dataset of 12,774 multivariate observations collected from 41 APMC mandis in "
        "Karnataka, India. Instead of applying the same predictive modeling approach to all structurally dissimilar commodities, we "
        "propose commodity-specific regressors tailored towards each commodity perishability and storage mechanism: OLS Linear "
        "Regression for Onion with high persistence and volatile Tomato, and Gradient Boosted Regression Trees for cold storage "
        "Potato. Predictive models are evaluated with respect to their chronological testing sets along with the benchmarking of "
        "Naive Persistence approach and Mean Absolute Percentage Error (MAPE). The system integrates the price forecasts into the "
        "working of Net Holding Return (NHR) mechanism involving an empirical risk buffer threshold (θ = Rs. 50/qtl) and a spatially "
        "intelligent market recommendation engine for maximizing the Net Realized Price by imposing penalties on distance-based mandis "
        "through regional transportation charges (c = Rs. 2.0/km/qtl). The web application for implementing the system has been built "
        "using the Streamlit framework."
    )
    r_abs_body.italic = True

    p_kw = add_p(space_after=8)
    r_kw_label = p_kw.add_run("Keywords— ")
    r_kw_label.bold = True
    r_kw_label.italic = True
    p_kw.add_run("Agricultural Price Forecasting, Decision Support System, Market Recommendation, Gradient Boosting, Linear Regression, Spatial Price Arbitrage, AGMARKNET.")

    # ==================== SECTION I ====================
    add_heading("I. INTRODUCTION")
    add_p(
        "Markets within the agricultural economy of India suffer from extensive structural fragmentation, information asymmetry, "
        "and volatility in prices. The farmers have to take decisions regarding the harvesting and selling of their crops amidst great "
        "uncertainty. Traditional sources of market information including official government portals like AGMARKNET operate as "
        "backward-looking information reporting agencies. They capture daily transaction rates at APMC markets, but they do not offer "
        "any forward-looking guidance or decision-making aid in regards to asset disposal after harvest. Therefore, small-holder "
        "farmers end up selling their crops immediately post-harvest due to market saturation and localized distress."
    )
    add_p(
        "Earlier work has utilized machine learning and deep sequence models for predicting agricultural spot prices. Yet, existing "
        "studies exhibit two primary weaknesses:"
    )
    p_c1 = add_p()
    p_c1.add_run("1. The Single-Model Fallacy: ").bold = True
    p_c1.add_run(
        "Researchers have frequently applied uniform model architectures (e.g., deep sequence predictors or generic ensemble trees) "
        "across commodities that inherently possess contrasting shelf-lives, storage capabilities, and arrival distributions."
    )
    p_c2 = add_p()
    p_c2.add_run("2. The Lack of Economic Perspective: ").bold = True
    p_c2.add_run(
        "Existing literature treats price forecasting purely as an isolated statistical curve-fitting exercise. Suggesting a market "
        "solely on the basis of the highest spot price without accounting for haulage overhead inevitably yields sub-optimal "
        "recommendations, frequently causing net financial losses once logistics costs are deducted."
    )
    add_p(
        "To address these operational deficits, this paper presents an applied, economically disciplined Decision Support System "
        "customized for the agrarian markets of Karnataka. The core contributions of this study are three-fold:"
    )
    bullet1 = add_p()
    bullet1.add_run("• Formulation of Heterogeneous Models: ").bold = True
    bullet1.add_run(
        "We demonstrate that post-harvest perishability dictates optimal model selection. Highly perishable produce (Tomato) and "
        "persistent commodities (Onion) are modeled using compact autoregressive OLS formulations, whereas commodities with non-linear "
        "cold-storage releases (Potato) require Gradient Boosted Decision Trees."
    )
    bullet2 = add_p()
    bullet2.add_run("• Spatial Logistics Optimization: ").bold = True
    bullet2.add_run(
        "Incorporation of road freight tariffs (c = Rs. 2.0/km/qtl) allows dynamic calculation of the Net Realized Price, eliminating "
        "deceptive spot-price arbitrage."
    )
    bullet3 = add_p()
    bullet3.add_run("• Risk-Buffered Sell/Wait Strategy: ").bold = True
    bullet3.add_run(
        "A Net Holding Return (NHR) engine equipped with an empirical θ = Rs. 50/qtl threshold shields farmers from false holding "
        "signals triggered by intraday market stochasticity."
    )

    # ==================== SECTION II ====================
    add_heading("II. RELATED WORK")
    add_p(
        "Quantitative agricultural price forecasting has evolved from classical univariate time-series to modern machine learning ensembles. "
        "Purohit et al. [1] investigated hybrid ARIMA-LSTM-SVM models within Indian markets, demonstrating that statistical-ML hybrids "
        "outperform single estimators. Sabu and Kumar [8] analyzed arecanut auctions in Kerala via statistical autoregression, though their "
        "approach remained bounded to a single regional perennial commodity."
    )
    add_p(
        "Singh and Sindhu [3] conducted comparative benchmarking across Decision Trees, Random Forests, and SVMs, deploying an elementary "
        "web-based interface. Manogna et al. [4] benchmarked deep RNNs (LSTM, GRU) against tree ensembles across 165 mandis and 23 crops, "
        "concluding that recurrent models thrive where deep uninterrupted sequence histories exist. Praveenkumar et al. [5] and Patil et al. [6] "
        "further confirmed the efficacy of attention-based recurrent networks and boosted ensembles on potato pricing and multi-commodity food "
        "security monitoring, respectively."
    )

    # --- Table I ---
    p_tlabel = add_p("TABLE I. COMPARATIVE TAXONOMY OF AGRICULTURAL FORECASTING LITERATURE", align=WD_ALIGN_PARAGRAPH.CENTER, bold=True)
    p_tlabel.paragraph_format.keep_with_next = True
    t1 = doc.add_table(rows=10, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER

    headers = ["Ref.", "Method", "Contribution", "Limitation"]
    row_data = [
        ("[1]", "Hybrid ARIMA, LSTM, SVM", "Hybrid price forecasting", "No decision support"),
        ("[2]", "LSTM", "Crop price prediction", "Prediction only"),
        ("[3]", "RF, SVM, DT", "Web-based prediction", "No market recommendation"),
        ("[4]", "LSTM, GRU, XGBoost", "Multi-model comparison", "Benchmarking only"),
        ("[5]", "LSTM, Bi-LSTM, AM-LSTM", "Potato price forecasting", "Single crop focus"),
        ("[6]", "LSTM, RF, XGBoost", "Food security forecasting", "Policy-level focus"),
        ("[7]", "Hybrid ARIMA-LSTM", "Productivity & forecasting", "No selling guidance"),
        ("[8]", "ARIMA, Regression", "Statistical forecasting", "Single crop & region"),
        ("Proposed", "OLS LR, Gradient Boosting", "Crop-specific forecast + Freight ranking + Risk NHR", "Three staple crops evaluated")
    ]

    for c_idx, h_text in enumerate(headers):
        c = t1.cell(0, c_idx)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h_text)
        r.bold = True
        r.font.size = Pt(8.5)

    for r_idx, row_vals in enumerate(row_data):
        for c_idx, val in enumerate(row_vals):
            c = t1.cell(r_idx + 1, c_idx)
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(val)
            r.font.size = Pt(8)
            if r_idx == len(row_data) - 1:
                r.bold = True

    add_p(
        "As summarized in Table I, existing literature predominantly frames market forecasting as an isolated predictive task. "
        "Practical considerations such as produce spoilage while holding, spatial logistics costs, and naive persistence benchmarking "
        "are consistently overlooked.", space_before=4
    )

    # ==================== SECTION III ====================
    add_heading("III. PROPOSED METHODOLOGY")
    add_p(
        "The proposed Decision Support System couples time-series machine learning models with an applied agricultural economic framework. "
        "Fig. 1 depicts the complete operational pipeline."
    )
    add_p("[Insert Fig. 1 Here: End-to-end architectural schematic of the freight-aware Decision Support System.]", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    add_subheading("A. Empirical Dataset and Preprocessing")
    add_p(
        "Daily auction records were acquired from the AGMARKNET portal across 41 regulated APMC markets in Karnataka covering 24 months "
        "(June 2023 to June 2025). The broader dataset was isolated into 12,774 cleaned records representing three primary staple commodities: "
        "Onion (Allium cepa), Tomato (Solanum lycopersicum), and Potato (Solanum tuberosum). Outliers above the 99th percentile and erroneous "
        "zero-price rows were removed. Missing auction trading days were filled via forward/backward interpolation for storable crops (Onion, Potato), "
        "while Tomato records remained on a discontinuous trading sequence to honor strict perishability."
    )

    add_subheading("B. Feature Engineering and Commodity Characterization")
    add_p(
        "Feature pipelines were engineered to align with commodity shelf-life dynamics: "
        "Perishables (Tomato) employ short-horizon autoregressive vectors: Lag1, Lag2, Lag3, and a 3-day moving average (MA3). "
        "Storable commodities (Onion, Potato) integrate weekly cyclical memory: Lag1, Lag2, Lag7, a 7-day moving average (MA7), and a seasonal monthly index (Mt)."
    )

    add_subheading("C. Model Training and Algorithmic Selection")
    add_p(
        "Three distinct algorithmic paradigms were benchmarked: Ordinary Least Squares (OLS) Linear Regression, Random Forest Regression "
        "(n_estimators=100, max_depth=10), and Gradient Boosted Decision Trees (GBDT: n_estimators=100, learning_rate=0.1, max_depth=3)."
    )

    add_subheading("D. Economic Decision and Mandi Recommendation Formulation")
    add_p(
        "The decision layer converts spot forecasts into financial directives using two modules:\n"
        "1) Net Holding Return (NHR): The expected return for holding stock to day t+1 is modeled as:\n"
        "     E[R_{t+1}] = P_hat_{t+1} - P_t - S_t - C_spoil\n"
        "To avoid speculative holding on noisy fluctuations, a risk buffer threshold is enforced:\n"
        "     Decision = 'WAIT' if E[R_{t+1}] >= θ else 'SELL'  (where θ = Rs. 50/qtl).\n"
        "2) Freight-Penalized Spatial Recommendation: Net Realized Price (NRP) is computed as:\n"
        "     NRP_m = P_m - (c × d_m)\n"
        "where d_m represents distance in kilometres, and c = Rs. 2.0/km/qtl based on regional LCV haulage rates."
    )

    # ==================== SECTION IV ====================
    add_heading("IV. IMPLEMENTATION")
    add_p(
        "The complete architecture is implemented in Python 3.11 with a modular Streamlit web interface. Predictive pipelines, scalers, "
        "and estimators are serialized via Joblib. The code base is decoupled into data ingestion (data_loader.py), inference execution "
        "(predict.py), spatial ranking (market.py), and user interaction (app.py)."
    )

    # ==================== SECTION V ====================
    add_heading("V. RESULTS AND DISCUSSION")
    add_subheading("A. Predictive Model Benchmarking")
    add_p(
        "Chronological 80/20 train/test splits were executed to avoid look-ahead bias. Candidate models were rigorously evaluated against "
        "the Naive Persistence baseline across RMSE, MAE, MAPE, and R2 metrics."
    )
    add_p("[Insert Fig. 2 Here: Comparative performance of candidate machine learning models against the Naive persistence baseline.]", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)
    add_p(
        "In the case of Onion, due to the high autocorrelation from one day to the next, the OLS regression model (R2=0.8587) and the "
        "persistence approach (R2=0.8584) yield near-equivalent explanatory power, confirming strong market price inertia. "
        "For Tomato, OLS regression boosts R2 to 0.3705 (+41.2% over persistence) and reduces MAPE to 18.94%. For Potato, Gradient Boosted "
        "Trees achieve superior performance (R2=0.3050 vs. 0.2090 naive baseline) by capturing non-linear cold-storage release dynamics."
    )

    add_subheading("B. Spatial Decision-Support Validation")
    add_p(
        "A practical case illustrates the necessity of freight penalization: Belgaum APMC offers a spot price of Rs. 1,800/qtl, whereas local "
        "Ramanagara APMC posts Rs. 1,650/qtl. Hauling 505 km to Belgaum incurs Rs. 1,010/qtl in freight, yielding a Net Realized Price of only "
        "Rs. 790/qtl. Conversely, transporting 45 km to Ramanagara incurs only Rs. 90/qtl, delivering Rs. 1,560/qtl net return—preventing a severe "
        "Rs. 770/qtl net loss for the cultivator."
    )
    add_p("[Insert Fig. 3 Here: Integrated decision-support output displaying price forecast, net profit, risk directive, and mandi ranking.]", align=WD_ALIGN_PARAGRAPH.CENTER, italic=True)

    # ==================== SECTION VI ====================
    add_heading("VI. CONCLUSION AND FUTURE WORK")
    add_p(
        "This research establishes an end-to-end Decision Support System bridging statistical commodity forecasting with post-harvest "
        "spatial economics. Future extensions will focus on automated AGMARKNET API ingestion, gridded satellite NDVI integrations, "
        "dynamic multi-farmer load consolidation, and vernacular Kannada voice/SMS delivery."
    )

    # ==================== REFERENCES ====================
    add_heading("REFERENCES")
    refs = [
        "[1] S. K. Purohit, S. Panigrahi, P. K. Sethy, and S. K. Behera, \"Time series forecasting of price of agricultural products using hybrid methods,\" Appl. Artif. Intell., vol. 35, no. 15, pp. 1388–1406, 2021.",
        "[2] D. Grewal and M. D. Daneshyari, \"Machine learning prediction of agricultural produces for Indian farmers using LSTM,\" Int. J. Multidiscip. Res. Growth Eval., vol. 3, no. 5, pp. 113–119, Sep.–Oct. 2022, doi: 10.54660/anfo.2022.3.5.5.",
        "[3] N. Singh and R. Sindhu, \"Crop price prediction using machine learning,\" J. Electr. Syst., vol. 20, no. 7s, pp. 2258–2269, 2024.",
        "[4] R. L. Manogna, V. Dharmaji, and S. Sarang, \"Enhancing agricultural commodity price forecasting with deep learning,\" Sci. Rep., vol. 15, Art. no. 20903, 2025, doi: 10.1038/s41598-025-05103-z.",
        "[5] A. Praveenkumar, G. K. Jha, S. D. Madival, A. Lama, and R. R. Kumar, \"Deep learning approaches for potato price forecasting: comparative analysis of LSTM, Bi-LSTM, and AM-LSTM models,\" Potato Res., vol. 68, pp. 1941–1963, 2025, doi: 10.1007/s11540-024-09823-z.",
        "[6] A. Patil, D. Shah, A. Shah, and R. Kotecha, \"Forecasting prices of agricultural commodities using machine learning for global food security: towards sustainable development goal 2,\" Int. J. Eng. Trends Technol., vol. 71, no. 12, pp. 277–291, Dec. 2023, doi: 10.14445/22315381/IJETT-V71I12P226.",
        "[7] M. Meeradevi, I. G. S. Yasaswi, M. R. Mundada, D. Sarika, and H. Shetty, \"Hybrid decision support system framework for enhancing crop productivity using machine learning,\" in Proc. 2nd Int. Conf. Recent Trends Mach. Learn., IoT, Smart Cities Appl., Lecture Notes Netw. Syst., vol. 237, Springer, 2022, pp. 57–66, doi: 10.1007/978-981-16-6407-6_6.",
        "[8] K. M. Sabu and T. K. Manoj Kumar, \"Predictive analytics in agriculture: Forecasting prices of arecanuts in Kerala,\" Procedia Comput. Sci., vol. 171, pp. 699–708, 2020, doi: 10.1016/j.procs.2020.04.076."
    ]

    for ref in refs:
        p_ref = add_p(ref, space_after=3, line_spacing=1.0)
        p_ref.paragraph_format.left_indent = Inches(0.25)
        p_ref.paragraph_format.first_line_indent = Inches(-0.25)
        p_ref.paragraph_format.keep_together = True

    # Save output
    output_filename = "final_ieee_manuscript.docx"
    doc.save(output_filename)
    print(f"Successfully generated: {output_filename}")

if __name__ == "__main__":
    create_document()